#!/usr/bin/env python3
"""
Strategy kernel example deck: REM Medical.

Reference layout for the kernel-of-good-strategy framework's deliverable.
Page one is the kernel; the pages behind it carry the detail. Edit this
script and re-run it; never hand-edit the .pptx.

    python3 build-example-deck.py

To render a real kernel, copy this script next to the strategy document and
replace everything below the "content" marker: the company name, the output
file name, and the locked lines. Nothing above that marker names the company.

Content: REM Medical was a sleep-medicine company that no longer exists.
Purpose, challenge, guiding policy, key actions, owners, and the year's
goals are from its original plan, used with its founder's permission. The
"why" behind the challenge and policy, the kill list, and the
assumptions and questions were not on the original page and are
illustrative: they show the format, not REM Medical's actual decisions.

Design: the Conscious OS design principles ("a well-set printed page with
one hot color"). Warm paper and warm ink, serif for what is said and sans
for what is controlled, hairline rules and cream bands instead of cards,
softened corners (rounder than the brand's rule, by choice), no shadows,
never a bold headline. Ember marks the
numerals; navy marks the guiding policy. Page one keeps the kernel's shape
(banner, arrow, block, arrow, three cards) in those surfaces. When rendering a real kernel,
replace the design tokens below with the project's design-principles.md
values; the layout does not change.

Typeface note: "New York" is the head of the serif stack and ships with
macOS; other machines fall back to Georgia.

Requires python-pptx (pip install python-pptx).
"""

import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- design tokens
# Values from the Conscious OS design-principles.md. Swap these for a
# project's own tokens; nothing below refers to a color by name.

PAPER = RGBColor(0xF7, 0xF2, 0xEA)      # cream-100, the page
CREAM_50 = RGBColor(0xFC, 0xFA, 0xF6)   # the lightest surface: page one's cards
CREAM_200 = RGBColor(0xEF, 0xE8, 0xDC)  # bands, the rare tinted panel
CREAM_300 = RGBColor(0xDF, 0xD6, 0xC8)  # hairlines that separate
BARK_900 = RGBColor(0x1A, 0x1A, 0x1A)   # headlines
BARK_700 = RGBColor(0x2D, 0x2D, 0x2D)   # reading text
BARK_500 = RGBColor(0x55, 0x55, 0x55)   # labels and metadata
BARK_400 = RGBColor(0x8F, 0x89, 0x81)   # non-text: rules that carry meaning
NAVY = RGBColor(0x1E, 0x28, 0x47)       # forest-500: competence, the guiding policy
NAVY_100 = RGBColor(0xE8, 0xEB, 0xF2)   # forest-100: the purpose banner's tint
EMBER = RGBColor(0xC2, 0x41, 0x0C)      # ember-500: heat, the numerals

# Corner rounding for bands and cards, in inches. The brand's own rule is
# near-square corners; this deck rounds them further by its author's choice.
RADIUS = 0.14

SERIF = "New York"        # what is said
SANS = "Helvetica Neue"   # what is controlled

MARGIN = 0.85
COLW = 11.63
FOOT_Y = 6.95

prs = Presentation()
prs.slide_width = Emu(12191695)
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- primitives


def flatten(sh):
    """Drop the theme style reference so no renderer adds a drop shadow."""
    st = sh.element.find(qn("p:style"))
    if st is not None:
        sh.element.remove(st)
    return sh


def slide(bg_color=PAPER):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = bg_color
    r.line.fill.background()
    r.shadow.inherit = False
    flatten(r)
    return s


def _shape(s, kind, x, y, w, h, fill):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    flatten(sh)
    return sh


def panel(s, x, y, w, h, fill=CREAM_200):
    """A band with softened corners: flat fill, no border, no shadow."""
    return card(s, x, y, w, h, fill)


def card(s, x, y, w, h, fill=CREAM_200, radius=RADIUS):
    """A band with softened corners: flat fill, no border, no shadow.
    radius is in inches, so a short banner and a tall card round the same."""
    sh = _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    return sh


def arrow_down(s, y, color=BARK_400):
    """A small triangle on the page's center line. Non-text, so BARK_400."""
    _shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 6.52, y, 0.30, 0.16, color).rotation = 180


def rule(s, x, y, w, color=CREAM_300):
    """Hairline. CREAM_300 separates; BARK_400 is a boundary that carries meaning."""
    return _shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.011, color)


def vrule(s, x, y, h, color=CREAM_300):
    return _shape(s, MSO_SHAPE.RECTANGLE, x, y, 0.011, h, color)


def text(s, x, y, w, h, runs, size=14, bold=False, color=BARK_700, font=SERIF,
         align=PP_ALIGN.LEFT, spacing=None, line_spacing=1.35, anchor=None, gap=0):
    """runs: a string, a list of (text, bold, color, size, font) tuples for one
    paragraph, or a list of those lists for several paragraphs."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor:
        tf.vertical_anchor = anchor
    if isinstance(runs, str):
        paras = [[(runs, bold, color)]]
    elif runs and isinstance(runs[0], tuple):
        paras = [runs]
    else:
        paras = runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(gap if i < len(paras) - 1 else 0)
        if isinstance(para, tuple):
            para = [para]
        for item in para:
            t = item[0]
            b = item[1] if len(item) > 1 and item[1] is not None else bold
            c = item[2] if len(item) > 2 and item[2] is not None else color
            sz = item[3] if len(item) > 3 and item[3] is not None else size
            f = item[4] if len(item) > 4 and item[4] is not None else font
            r = p.add_run()
            r.text = t
            r.font.name = f
            r.font.size = Pt(sz)
            r.font.bold = b
            r.font.color.rgb = c
            if spacing:
                r.font._rPr.set("spc", str(int(spacing)))
    return box


def label(s, txt, x, y, w=4.0, color=BARK_500, size=11):
    """Small-caps sans label: uppercase, 0.12em tracking, never bold."""
    return text(s, x, y, w, 0.18, txt.upper(), size=size, color=color, font=SANS,
                spacing=132, line_spacing=1.0)


def owner_run(name, size=10):
    """An owner's name as a small-caps sans run, for the end of a serif line."""
    return ("   " + name.upper(), False, BARK_500, size, SANS)


def eyebrow(s, txt, y=0.52):
    return label(s, txt, MARGIN, y, COLW)


def h1(s, runs, y=0.86, size=28, w=COLW, h=0.50, line_spacing=1.14, color=BARK_900):
    """Never bold a headline. Authority comes from size and space."""
    return text(s, MARGIN, y, w, h, runs, size=size, color=color, font=SERIF,
                line_spacing=line_spacing)


def numbered_h1(s, n, txt, y=0.86, size=28, h=0.50):
    """Section opener: a numeral in ember, then the serif heading."""
    return h1(s, [(n, False, EMBER), ("   " + txt, False, BARK_900)], y=y, size=size, h=h)


def footer(s, page):
    label(s, DECK_NAME, MARGIN, FOOT_Y, 10.9, size=9)
    text(s, 11.9, FOOT_Y, 0.60, 0.16, str(page), size=9, color=BARK_500, font=SANS,
         align=PP_ALIGN.RIGHT, line_spacing=1.0)


def est_h(body, size, w, line_spacing=1.35, gap_pt=0):
    """Estimated height, in inches, of a string or a list of strings wrapped in a box
    w inches wide at size pt. python-pptx cannot measure text, so the pages flow on
    this estimate: about 0.55 em per character, plus 15% for ragged wrapping. It
    runs a little long on purpose; a column that steps down a point is better than
    one whose panel lands on its last line."""
    paras = body if isinstance(body, list) else [body]
    cpl = max(1, int(w * 72 / (0.55 * size)))
    n = sum(max(1, math.ceil(len(t) * 1.15 / cpl)) for t in paras)
    # A rendered line is the font's natural height, about 1.17 em, times the spacing.
    return n * size * 1.17 * line_spacing / 72 + (len(paras) - 1) * gap_pt / 72


def labelled_text(s, x, y, w, lbl, body, size=13.5, color=BARK_700, label_color=BARK_500):
    """Label, then a body sized to its text. Returns the y where the next block starts."""
    label(s, lbl, x, y, w, label_color)
    h = est_h(body, size, w)
    text(s, x, y + 0.26, w, h, body, size=size, color=color)
    return y + 0.26 + h + 0.22


def labelled_panel(s, x, y, w, lbl, body, size=13.5, pad=0.28):
    """A tinted panel sized to its text: the closing block of a column."""
    h = est_h(body, size, w - 2 * pad) + 0.74
    panel(s, x, y, w, h)
    label(s, lbl, x + pad, y + 0.24, w - 2 * pad)
    text(s, x + pad, y + 0.50, w - 2 * pad, h - 0.62, body, size=size, color=BARK_700)
    return y


def lines(s, x, y, w, items, size=13.5, color=BARK_700, gap=7):
    """Short serif lines, one paragraph per item, no bullet glyphs. They flow in one
    box sized to the estimate, so a line that wraps pushes the next one down instead
    of into it. Returns the y where the next block starts."""
    paras = [[(item, False, color)] if isinstance(item, str) else item for item in items]
    plain = ["".join(run[0] for run in para) for para in paras]
    h = est_h(plain, size, w, line_spacing=1.3, gap_pt=gap)
    text(s, x, y, w, h, paras, size=size, color=color, line_spacing=1.3, gap=gap)
    return y + h + 0.22


def labelled_lines(s, x, y, w, lbl, items, size=13.5):
    label(s, lbl, x, y, w)
    return lines(s, x, y + 0.26, w, items, size=size)


def column(s, x, top, w, blocks, closing, bottom):
    """A detail-page column: labelled blocks in order, then a tinted closing panel.
    blocks: ("text", label, body, size, color, label_color) or ("lines", label, items, size).
    closing: (label, body). The column is measured first; if it would run past
    bottom, every size steps down together until it fits. The panel sits at the
    foot of the column when there is room, and right after the last block when
    there isn't."""
    for scale in (1.0, 0.94, 0.88, 0.82, 0.76, 0.70, 0.65):
        y = top
        for b in blocks:
            if b[0] == "text":
                y += 0.26 + est_h(b[2], b[3] * scale, w) + 0.22
            else:
                plain = ["".join(r[0] for r in it) if isinstance(it, list) else it for it in b[2]]
                y += 0.26 + est_h(plain, b[3] * scale, w, 1.3, 7) + 0.22
        ph = est_h(closing[1], 13.5 * scale, w - 0.56) + 0.74
        if y + ph <= bottom:
            break
    y = top
    for b in blocks:
        if b[0] == "text":
            y = labelled_text(s, x, y, w, b[1], b[2], size=b[3] * scale, color=b[4], label_color=b[5])
        else:
            y = labelled_lines(s, x, y, w, b[1], b[2], size=b[3] * scale)
    labelled_panel(s, x, max(y, bottom - ph), w, closing[0], closing[1], size=13.5 * scale)


# ---------------------------------------------------------------- content
# Everything below this line is the strategy. Replace all of it.

COMPANY = "REM Medical"
DECK_NAME = f"{COMPANY} strategy kernel"
OUT = Path(__file__).with_name("example-strategy-kernel.pptx")

PURPOSE = "Enable people to lead a happier, healthier life starting with a better night's sleep"

CHALLENGE = "It is likely that Medicare reimbursement for sleep diagnostics will be significantly reduced."
CHALLENGE_WHY = ("Medicare has signaled cuts to in-lab polysomnography and is steering diagnosis toward "
                 "home sleep testing at a fraction of the rate. Commercial payers follow Medicare within "
                 "a year or two. Our revenue is built on the in-lab test.")
CHALLENGE_WRONG_IF = ("The next Medicare fee schedule holds in-lab sleep study rates flat, and commercial "
                      "payers don't move.")

POLICY = ("Leverage our strength in delivering outcomes by shifting from fee for service "
          "to getting paid for results.")
POLICY_WHY = [
    "Our outcomes data is the one asset a rate cut can't touch",
    "Payers cutting the test will still pay for a treated, adherent patient",
    "Fee-for-service funds the transition for as long as it lasts",
]
POLICY_GATE = ("When something new arrives, ask one question: does it get us paid for results, or fund "
               "the engine that will? If neither, it goes on the kill list.")

ACTIONS = [
    {
        "n": "1",
        "title": "Fuel the engine",
        "owner": "Jeff",
        "what": ("The existing fee-for-service business, in-lab and home testing, pays for the "
                 "shift. Keep it full and keep it collecting while reimbursement still holds."),
        "goals": [
            ("Engage 700 patients a month in Q4 in the Arizona region", "Jeff"),
            ("Perform, and collect payment on, 250 home tests in Q4", "Robert"),
        ],
        "depends": [
            "Referral volume from the existing physician network holds through the year",
            "Home-test billing and collection run without a new hire",
        ],
        "signal": "Monthly engaged patients and home-test collections, reviewed on the first of the month",
    },
    {
        "n": "2",
        "title": "New delivery models",
        "owner": "Eric",
        "what": ("Extend into new ways to deliver care. Find buyers who pay for a treated, sleeping "
                 "patient rather than a test: employers through occupational health, patients through "
                 "ancillary products, sponsors through research."),
        "goals": [
            ("Treat 2,000 patients through occupational health", "Jen F"),
            ("$100,000 a month from ancillary DME products by end of Q3", "Michelle"),
            ("First dollar from the research division by end of Q3", "Colleen"),
        ],
        "depends": [
            "An occupational-health contract that pays per treated employee, not per test",
            "DME margin that survives payers following Medicare",
        ],
        "signal": "Revenue by delivery model, with fee-for-service share falling quarter over quarter",
    },
    {
        "n": "3",
        "title": "Geographic reach",
        "owner": "Russell",
        "what": ("Extend into new regions where sleep care is thin. Build each site with provider "
                 "groups who bring the patients, so it opens with referrals already attached."),
        "goals": [
            ("See the first patient in the Pacific Northwest by Aug 31", "Ken"),
            ("Break ground on 4 new sites in Q4", "Jen B"),
            ("Sign 6 provider groups to co-build new facilities", "Russell"),
        ],
        "depends": [
            "Provider groups willing to co-invest rather than refer at arm's length",
            "Site economics that work at reduced in-lab rates",
        ],
        "signal": "Signed co-build contracts and days from signature to first patient",
    },
]

KILL_LIST = [
    ("In-lab beds in existing markets", "In-lab utilization above 85% for two straight quarters"),
    ("Price cuts to win referrals", "A referral source names price as the reason they left"),
    ("Consumer sleep products", "Two occupational-health clients ask for them"),
    ("Acquiring other sleep labs", "A results-based payer contract is signed and needs capacity"),
]

ASSUMPTIONS = [
    ("Payers will pay for outcomes, not only for tests",
     "Two payer conversations end with “we only pay per study”"),
    ("Home testing can be profitable at reduced rates",
     "Home-test margin is negative after the Q4 volume push"),
    ("Occupational-health buyers want treated employees, not diagnoses",
     "Pilots renew as diagnostic-only contracts"),
]

QUESTIONS = [
    ("What does a results-based contract look like to a regional payer?", "Eric"),
    ("Which Pacific Northwest sites have the referral density to support a lab?", "Russell"),
    ("What DME margin holds once commercial payers follow Medicare?", "Michelle"),
]

# Two-column geometry shared by the detail pages.
COL_W = 5.55
X2 = MARGIN + COL_W + 0.53
VR_X = MARGIN + COL_W + 0.26
BODY_TOP = 1.66          # first label on a detail page
BODY_BOTTOM = 6.60       # lowest edge content may reach; the footer sits at FOOT_Y

# ---------------------------------------------------------------- 1. the kernel
# Page one keeps the kernel's shape: a purpose banner, an arrow down to the
# challenge-and-policy block, an arrow down to three key-action cards. The
# brand supplies the surfaces (navy tint, cream band, cream-50 cards), the
# type (serif for what is said), and the ember numerals. The bands grow with
# their sentences; the cards take whatever height is left above the footer.

s = slide()
eyebrow(s, "The kernel", y=0.34)
h1(s, DECK_NAME, y=0.56, size=22, h=0.36)

SENT_W = COLW - 2.38     # width of the sentence column inside a band

# The bands take the height their sentences need, at the largest size that
# still leaves the cards room: the cards must start by 3.8in.
for purpose_sz, sent_sz in ((15, 14), (14, 13), (13, 12), (12, 11)):
    ph = max(0.60, est_h(PURPOSE, purpose_sz, SENT_W, 1.1) + 0.24)
    ch = max(0.56, est_h(CHALLENGE, sent_sz, SENT_W, 1.2) + 0.20)
    pol = max(0.56, est_h(POLICY, sent_sz, SENT_W, 1.2) + 0.20)
    if 1.04 + ph + 0.32 + ch + pol + 0.32 <= 3.8:
        break

# Purpose banner
y = 1.04
card(s, MARGIN, y, COLW, ph, NAVY_100)
label(s, "Purpose", MARGIN + 0.28, y + 0.21, 1.6, NAVY, size=10)
text(s, MARGIN + 2.10, y + 0.08, SENT_W, ph - 0.16, PURPOSE, size=purpose_sz, color=NAVY,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
y += ph + 0.08
arrow_down(s, y)
y += 0.24

# Challenge and guiding policy, one band, a hairline between
card(s, MARGIN, y, COLW, ch + pol, CREAM_200)
label(s, "Challenge", MARGIN + 0.28, y + ch / 2 - 0.08, 2.0, size=10)
text(s, MARGIN + 2.10, y + 0.08, SENT_W, ch - 0.16, CHALLENGE, size=sent_sz, color=BARK_900,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)
rule(s, MARGIN + 0.28, y + ch, COLW - 0.56, BARK_400)
label(s, "Guiding policy", MARGIN + 0.28, y + ch + pol / 2 - 0.08, 2.4, NAVY, size=10)
text(s, MARGIN + 2.10, y + ch + 0.08, SENT_W, pol - 0.16, POLICY, size=sent_sz, color=NAVY,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)
y += ch + pol + 0.08
arrow_down(s, y)
y += 0.24

# Three key-action cards
CARD_Y, GAP = y, 0.30
CARD_H = 6.66 - CARD_Y
CW = (COLW - 2 * GAP) / 3
PAD = 0.24
for i, a in enumerate(ACTIONS):
    x = MARGIN + i * (CW + GAP)
    card(s, x, CARD_Y, CW, CARD_H, CREAM_50)
    text(s, x + PAD, CARD_Y + 0.26, CW - 2 * PAD, 0.32,
         [(a["n"], False, EMBER, 16, SERIF), ("   " + a["title"], False, BARK_900, 16, SERIF)],
         align=PP_ALIGN.CENTER, line_spacing=1.1)
    text(s, x + PAD, CARD_Y + 0.64, CW - 2 * PAD, 0.2, "Owner: " + a["owner"], size=10.5,
         color=BARK_500, font=SANS, align=PP_ALIGN.CENTER, line_spacing=1.2)
    rule(s, x + PAD, CARD_Y + 0.94, CW - 2 * PAD)
    label(s, "This year", x + PAD, CARD_Y + 1.08, 1.5, size=9.5)
    # Goals shrink a point when they would not fit the card.
    goal_size = 10.5
    goal_paras = [f"{goal}   {owner.upper()}" for goal, owner in a["goals"]]
    while goal_size > 8 and est_h(goal_paras, goal_size, CW - 2 * PAD, 1.3, 6) > CARD_H - 1.5:
        goal_size -= 0.5
    lines(s, x + PAD, CARD_Y + 1.32, CW - 2 * PAD,
          [[(goal, False, BARK_700, goal_size, SANS), owner_run(owner, goal_size - 1.5)]
           for goal, owner in a["goals"]],
          size=goal_size, gap=6)
footer(s, 1)

# ---------------------------------------------------------------- 2. challenge and policy in detail

s = slide()
eyebrow(s, "The challenge and the policy")
h1(s, "What the kernel stands on", y=0.82, size=26, h=0.44)
rule(s, MARGIN, 1.42, COLW, BARK_400)

column(s, MARGIN, BODY_TOP, COL_W, [
    ("text", "Challenge", CHALLENGE, 17, BARK_900, BARK_500),
    ("text", "Why it is that way", CHALLENGE_WHY, 13.5, BARK_700, BARK_500),
], ("How we'd know this is wrong", CHALLENGE_WRONG_IF), BODY_BOTTOM)

vrule(s, VR_X, BODY_TOP, BODY_BOTTOM - BODY_TOP)

column(s, X2, BODY_TOP, COL_W, [
    ("text", "Guiding policy", POLICY, 17, NAVY, NAVY),
    ("lines", "Why this policy fits", POLICY_WHY, 13.5),
], ("The gate", POLICY_GATE), BODY_BOTTOM)
footer(s, 2)

# ---------------------------------------------------------------- 3-5. one page per key action

for i, a in enumerate(ACTIONS):
    s = slide()
    eyebrow(s, f"Key action {a['n']}")
    numbered_h1(s, a["n"], a["title"], y=0.82, size=26, h=0.44)
    text(s, MARGIN, 1.34, COLW, 0.2, "Owner: " + a["owner"], size=11, color=BARK_500, font=SANS,
         line_spacing=1.0)
    rule(s, MARGIN, 1.66, COLW, BARK_400)
    top = 1.90

    column(s, MARGIN, top, COL_W, [
        ("text", "What it is", a["what"], 15, BARK_900, BARK_500),
        ("lines", "What it depends on", a["depends"], 13.5),
    ], ("How we'll know it's working", a["signal"]), BODY_BOTTOM)

    vrule(s, VR_X, top, BODY_BOTTOM - top)

    label(s, "This year's goals", X2, top, COL_W)
    y = top + 0.30
    for goal, owner in a["goals"]:
        # One line per goal, the owner in small caps at its end, a hairline beneath.
        h = est_h(goal + "   " + owner, 15, COL_W, 1.3)
        text(s, X2, y, COL_W, h, [(goal, False, BARK_900, 15, SERIF), owner_run(owner)],
             line_spacing=1.3)
        rule(s, X2, y + h + 0.12, COL_W)
        y += h + 0.30
    footer(s, 3 + i)

# ---------------------------------------------------------------- 6. the kill list (its own page; not on page one)

s = slide()
eyebrow(s, "The kill list")
h1(s, "Not now, and what would bring each one back", y=0.82, size=26, h=0.44)
text(s, MARGIN, 1.40, 10.6, 0.5,
     "A refusal without a re-entry condition gets reopened in a hallway six months later. "
     "Each line here closes a door and puts a condition on the handle.",
     size=14, color=BARK_700)

NAME_W, COND_X = 5.0, 6.40
COND_W = MARGIN + COLW - COND_X
label(s, "Not now", MARGIN, 2.30, 4.5)
label(s, "Comes back when", COND_X, 2.30, 5.5)
rule(s, MARGIN, 2.58, COLW, BARK_400)
KILL_TOP = 2.78
# Four rows sit at 16/14; seven step down until every row fits above the footer.
for name_sz, cond_sz, gap in ((16, 14, 0.30), (15, 13, 0.24), (14, 12, 0.20), (13, 11, 0.16),
                              (12, 10.5, 0.14)):
    heights = [max(est_h(n, name_sz, NAME_W - 0.4, 1.25), est_h(c, cond_sz, COND_W, 1.3)) + gap
               for n, c in KILL_LIST]
    if sum(heights) <= BODY_BOTTOM - KILL_TOP:
        break
spare = max(0.0, BODY_BOTTOM - KILL_TOP - sum(heights)) / max(1, len(KILL_LIST))
y = KILL_TOP
for (name, cond), h in zip(KILL_LIST, heights):
    step = h + spare
    text(s, MARGIN, y, NAME_W, step - gap, [("×  ", False, BARK_400), (name, False, BARK_900)],
         size=name_sz, line_spacing=1.25)
    text(s, COND_X, y + 0.02, COND_W, step - gap, cond, size=cond_sz, color=BARK_700, line_spacing=1.3)
    rule(s, MARGIN, y + step - 0.16, COLW)
    y += step
footer(s, 6)

# ---------------------------------------------------------------- 7. assumptions and questions
# Two columns when both lists fit at a readable size. When they don't (five
# assumptions with two-line claims, say), each list gets a full-width page of
# its own instead of shrinking past the point of reading.

ROWS_TOP = 2.00


def row_heights(rows, first_sz, second_sz, gap, w):
    """Height of each (first line, second line) row at these sizes, gap included."""
    return [est_h(a, first_sz, w, 1.3) + est_h(b, second_sz, w, 1.3) + 0.08 + gap for a, b in rows]


def fit(rows, w):
    """The largest type on the ladder at which every row fits above the footer.
    Returns (first size, second size, gap, row heights, fits)."""
    ladder = ((15, 13, 0.30), (14, 12, 0.26), (13, 11.5, 0.22), (12, 11, 0.20))
    for first_sz, second_sz, gap in ladder:
        heights = row_heights(rows, first_sz, second_sz, gap, w)
        if sum(heights) <= BODY_BOTTOM - ROWS_TOP:
            spare = (BODY_BOTTOM - ROWS_TOP - sum(heights)) / max(1, len(rows))
            return first_sz, second_sz, gap, [h + spare for h in heights], True
    return first_sz, second_sz, gap, heights, False


def rows_column(s, x, w, rows, first_sz, second_sz, gap, heights, second_style):
    """Rows of two lines each with a hairline beneath. second_style is "broke" for
    an assumption's break condition or "owner" for a question's owner."""
    y = ROWS_TOP
    for (first, second), step in zip(rows, heights):
        if second_style == "broke":
            runs = [[(first, False, BARK_900, first_sz)],
                    [("BROKE WHEN   ", False, BARK_500, 9.5, SANS), (second, False, BARK_700, second_sz)]]
        else:
            runs = [[(first, False, BARK_900, first_sz)], [(second.upper(), False, BARK_500, 9.5, SANS)]]
        text(s, x, y, w, step - gap, runs, line_spacing=1.3, gap=6)
        rule(s, x, y + step - 0.16, w)
        y += step


a_rows = [(c, b) for c, b in ASSUMPTIONS]
q_rows = [(q, o) for q, o in QUESTIONS]
a_fit = fit(a_rows, COL_W)
q_fit = fit(q_rows, COL_W)

if a_fit[4] and q_fit[4]:
    s = slide()
    eyebrow(s, "WHAT WE'RE STANDING ON")
    h1(s, "Key assumptions and key questions", y=0.82, size=26, h=0.44)
    rule(s, MARGIN, 1.42, COLW, BARK_400)
    label(s, "Key assumptions", MARGIN, BODY_TOP, COL_W)
    rows_column(s, MARGIN, COL_W, a_rows, *a_fit[:4], "broke")
    vrule(s, VR_X, BODY_TOP, BODY_BOTTOM - BODY_TOP)
    label(s, "Key questions, each with an owner", X2, BODY_TOP, COL_W)
    rows_column(s, X2, COL_W, q_rows, *q_fit[:4], "owner")
    footer(s, 7)
else:
    s = slide()
    eyebrow(s, "WHAT WE'RE STANDING ON")
    h1(s, "Key assumptions", y=0.82, size=26, h=0.44)
    rule(s, MARGIN, 1.42, COLW, BARK_400)
    label(s, "What must be true, and how we'd know it broke", MARGIN, BODY_TOP, COLW)
    rows_column(s, MARGIN, COLW, a_rows, *fit(a_rows, COLW)[:4], "broke")
    footer(s, 7)

    s = slide()
    eyebrow(s, "WHAT WE'RE STANDING ON")
    h1(s, "Key questions", y=0.82, size=26, h=0.44)
    rule(s, MARGIN, 1.42, COLW, BARK_400)
    label(s, "What we don't know yet, and who owns finding out", MARGIN, BODY_TOP, COLW)
    rows_column(s, MARGIN, COLW, q_rows, *fit(q_rows, COLW)[:4], "owner")
    footer(s, 8)

prs.save(OUT)
print(f"wrote {OUT}")
